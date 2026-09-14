from __future__ import annotations

import os
import shutil
import sqlite3
from datetime import datetime
from decimal import Decimal, ROUND_FLOOR
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import (
    BACKUP_DIR,
    DB_PATH,
    EXPORT_DIR,
    MULTIPLO_REDONDEO,
    PORCENTAJE_CHOFER,
    PORCENTAJE_PROPIETARIO,
    TIPO_INGRESO_OTRO,
    TIPO_INGRESO_TRICICLO,
)
from database import conectar, inicializar_base_de_datos, transaccion
from utils import decimal_a_float


def _validar_monto_entero(
    monto: Decimal,
    permitir_cero: bool = False,
) -> Decimal:
    """Valida que un monto sea finito, entero y no negativo."""
    if not monto.is_finite():
        raise ValueError("El monto debe ser un número finito.")

    monto_entero = monto.to_integral_value()
    if monto != monto_entero:
        raise ValueError("El monto debe ser un número entero, sin centavos.")

    if monto < 0 or (not permitir_cero and monto == 0):
        raise ValueError("El monto debe ser mayor que cero.")

    return monto_entero


def calcular_reparto(
    monto_total: Decimal,
    balance_inicial: bool = False,
) -> dict[str, Decimal]:
    """
    Conserva exactamente la regla original:
    - Chofer: 25 %, redondeado hacia abajo al múltiplo de 50.
    - Propietario: 25 %, redondeado hacia abajo al múltiplo de 50.
    - Triciclo: el resto.
    - Balance inicial: 100 % propietario.
    """
    monto_total = _validar_monto_entero(monto_total, permitir_cero=True)

    if balance_inicial:
        return {
            "chofer": Decimal("0.00"),
            "propietario": monto_total,
            "triciclo": Decimal("0.00"),
        }

    base_chofer = monto_total * PORCENTAJE_CHOFER
    parte_chofer = (
        (base_chofer / MULTIPLO_REDONDEO).to_integral_value(rounding=ROUND_FLOOR)
        * MULTIPLO_REDONDEO
    )

    base_propietario = monto_total * PORCENTAJE_PROPIETARIO
    parte_propietario = (
        (base_propietario / MULTIPLO_REDONDEO).to_integral_value(rounding=ROUND_FLOOR)
        * MULTIPLO_REDONDEO
    )

    parte_triciclo = monto_total - parte_chofer - parte_propietario

    return {
        "chofer": parte_chofer.quantize(Decimal("0.01")),
        "propietario": parte_propietario.quantize(Decimal("0.01")),
        "triciclo": parte_triciclo.quantize(Decimal("0.01")),
    }


def existe_balance_inicial() -> bool:
    with conectar() as conn:
        fila = conn.execute(
            """
            SELECT 1
            FROM ingresos
            WHERE Es_Balance_Inicial = 1 OR Conductor = 'BALANCE INICIAL'
            LIMIT 1
            """
        ).fetchone()
        return fila is not None


def registrar_balance_inicial(fecha: str, monto: Decimal) -> int:
    monto = _validar_monto_entero(monto, permitir_cero=True)
    reparto = calcular_reparto(monto, balance_inicial=True)
    with transaccion() as conn:
        cursor = conn.execute(
            """
            INSERT INTO ingresos (
                Fecha, Conductor, Ingreso_Total, Parte_Triciclo,
                Parte_Chofer, Parte_Propietario, Es_Balance_Inicial, Nota
            )
            VALUES (?, 'BALANCE INICIAL', ?, ?, ?, ?, 1, 'Saldo inicial')
            """,
            (
                fecha,
                decimal_a_float(monto),
                decimal_a_float(reparto["triciclo"]),
                decimal_a_float(reparto["chofer"]),
                decimal_a_float(reparto["propietario"]),
            ),
        )
        return int(cursor.lastrowid)


def registrar_ingreso(
    fecha: str,
    conductor: str,
    monto: Decimal,
    nota: str = "",
) -> int:
    if not conductor.strip():
        raise ValueError("Selecciona un conductor.")

    monto = _validar_monto_entero(monto)
    reparto = calcular_reparto(monto)
    with transaccion() as conn:
        cursor = conn.execute(
            """
            INSERT INTO ingresos (
                Fecha, Conductor, Ingreso_Total, Parte_Triciclo,
                Parte_Chofer, Parte_Propietario, Parte_Otro, Tipo_Ingreso,
                Es_Balance_Inicial, Nota
            )
            VALUES (?, ?, ?, ?, ?, ?, 0, ?, 0, ?)
            """,
            (
                fecha,
                conductor.strip(),
                decimal_a_float(monto),
                decimal_a_float(reparto["triciclo"]),
                decimal_a_float(reparto["chofer"]),
                decimal_a_float(reparto["propietario"]),
                TIPO_INGRESO_TRICICLO,
                nota.strip(),
            ),
        )
        return int(cursor.lastrowid)


def registrar_otro_ingreso(
    fecha: str,
    monto: Decimal,
    nota: str = "",
) -> int:
    """Registra un ingreso externo sin repartirlo entre conductor y propietario."""
    monto = _validar_monto_entero(monto)
    with transaccion() as conn:
        cursor = conn.execute(
            """
            INSERT INTO ingresos (
                Fecha, Conductor, Ingreso_Total, Parte_Triciclo,
                Parte_Chofer, Parte_Propietario, Parte_Otro, Tipo_Ingreso,
                Es_Balance_Inicial, Nota
            )
            VALUES (?, '', ?, 0, 0, 0, ?, ?, 0, ?)
            """,
            (
                fecha,
                decimal_a_float(monto),
                decimal_a_float(monto),
                TIPO_INGRESO_OTRO,
                nota.strip(),
            ),
        )
        return int(cursor.lastrowid)


def registrar_gasto(
    fecha: str,
    monto: Decimal,
    categoria: str,
    comentario: str = "",
) -> int:
    monto = _validar_monto_entero(monto)
    with transaccion() as conn:
        cursor = conn.execute(
            """
            INSERT INTO gastos (Fecha, Monto, Categoria, Comentario)
            VALUES (?, ?, ?, ?)
            """,
            (
                fecha,
                decimal_a_float(monto),
                categoria.strip() or "General",
                comentario.strip(),
            ),
        )
        return int(cursor.lastrowid)


def obtener_conductores(solo_activos: bool = True) -> list[str]:
    consulta = "SELECT Nombre FROM conductores"
    parametros: tuple[Any, ...] = ()
    if solo_activos:
        consulta += " WHERE Activo = 1"
    consulta += " ORDER BY Nombre COLLATE NOCASE"

    with conectar() as conn:
        return [fila["Nombre"] for fila in conn.execute(consulta, parametros)]


def agregar_conductor(nombre: str) -> None:
    limpio = nombre.strip()
    if not limpio:
        raise ValueError("Escribe el nombre del conductor.")

    with transaccion() as conn:
        existente = conn.execute(
            "SELECT id FROM conductores WHERE Nombre = ? COLLATE NOCASE",
            (limpio,),
        ).fetchone()

        if existente:
            conn.execute(
                "UPDATE conductores SET Activo = 1, Nombre = ? WHERE id = ?",
                (limpio, existente["id"]),
            )
        else:
            conn.execute(
                "INSERT INTO conductores (Nombre, Activo) VALUES (?, 1)",
                (limpio,),
            )


def cambiar_estado_conductor(nombre: str, activo: bool) -> None:
    with transaccion() as conn:
        conn.execute(
            "UPDATE conductores SET Activo = ? WHERE Nombre = ?",
            (1 if activo else 0, nombre),
        )


def obtener_estadisticas_globales() -> dict[str, Any]:
    with conectar() as conn:
        totales = conn.execute(
            """
            SELECT
                COALESCE(SUM(Ingreso_Total), 0) AS total_ingresado,
                COALESCE(SUM(Parte_Triciclo), 0) AS total_triciclo,
                COALESCE(SUM(Parte_Propietario), 0) AS total_propietario,
                COALESCE(SUM(Parte_Otro), 0) AS total_otros,
                COALESCE(SUM(Parte_Chofer), 0) AS total_chofer
            FROM ingresos
            """
        ).fetchone()

        total_gastos = conn.execute(
            "SELECT COALESCE(SUM(Monto), 0) AS total FROM gastos"
        ).fetchone()["total"]

        salarios = {
            fila["Conductor"]: fila["total"]
            for fila in conn.execute(
                """
                SELECT Conductor, COALESCE(SUM(Parte_Chofer), 0) AS total
                FROM ingresos
                WHERE Es_Balance_Inicial = 0
                  AND Tipo_Ingreso = 'TRICICLO'
                GROUP BY Conductor
                ORDER BY Conductor
                """
            )
        }

    balance = (
        float(totales["total_triciclo"])
        + float(totales["total_propietario"])
        + float(totales["total_otros"])
        - float(total_gastos)
    )

    return {
        "total_ingresado": float(totales["total_ingresado"]),
        "total_triciclo": float(totales["total_triciclo"]),
        "total_propietario": float(totales["total_propietario"]),
        "total_otros": float(totales["total_otros"]),
        "total_chofer": float(totales["total_chofer"]),
        "total_gastos": float(total_gastos),
        "balance": balance,
        "salarios": salarios,
    }


def obtener_estadisticas_mes(anio: int, mes: int) -> dict[str, float]:
    patron = f"{anio}-{mes:02d}-%"
    with conectar() as conn:
        ingreso = conn.execute(
            """
            SELECT
                COALESCE(SUM(Ingreso_Total), 0) AS total,
                COALESCE(
                    SUM(Parte_Triciclo + Parte_Propietario + Parte_Otro), 0
                ) AS disponible,
                COALESCE(SUM(Parte_Otro), 0) AS otros
            FROM ingresos
            WHERE Fecha LIKE ?
            """,
            (patron,),
        ).fetchone()
        gastos = conn.execute(
            "SELECT COALESCE(SUM(Monto), 0) AS total FROM gastos WHERE Fecha LIKE ?",
            (patron,),
        ).fetchone()["total"]

    return {
        "ingresos": float(ingreso["total"]),
        "disponible_generado": float(ingreso["disponible"]),
        "otros": float(ingreso["otros"]),
        "gastos": float(gastos),
        "resultado": float(ingreso["disponible"]) - float(gastos),
    }


def obtener_ganancia_conductor(
    conductor: str,
    anio: int,
    mes: int,
) -> dict[str, float | int]:
    """Devuelve el salario asignado a un conductor en un mes concreto."""
    if not conductor:
        return {"salario": 0.0, "ingreso_bruto": 0.0, "registros": 0}

    patron = f"{anio}-{mes:02d}-%"
    with conectar() as conn:
        fila = conn.execute(
            """
            SELECT
                COALESCE(SUM(Parte_Chofer), 0) AS salario,
                COALESCE(SUM(Ingreso_Total), 0) AS ingreso_bruto,
                COUNT(*) AS registros
            FROM ingresos
            WHERE Conductor = ?
              AND Fecha LIKE ?
              AND Es_Balance_Inicial = 0
            """,
            (conductor, patron),
        ).fetchone()

    return {
        "salario": float(fila["salario"]),
        "ingreso_bruto": float(fila["ingreso_bruto"]),
        "registros": int(fila["registros"]),
    }


def obtener_movimientos(
    anio: int | None = None,
    mes: int | None = None,
    tipo: str = "Todos",
    conductor: str = "Todos",
    busqueda: str = "",
    limite: int | None = None,
    ascendente: bool = False,
) -> list[dict[str, Any]]:
    condiciones_ing: list[str] = []
    condiciones_gas: list[str] = []
    params_ing: list[Any] = []
    params_gas: list[Any] = []

    if anio is not None:
        if mes is not None:
            patron = f"{anio}-{mes:02d}-%"
        else:
            patron = f"{anio}-%"
        condiciones_ing.append("Fecha LIKE ?")
        condiciones_gas.append("Fecha LIKE ?")
        params_ing.append(patron)
        params_gas.append(patron)

    if conductor and conductor != "Todos":
        condiciones_ing.append("Conductor = ?")
        params_ing.append(conductor)

    texto = busqueda.strip()
    if texto:
        condiciones_ing.append(
            "(Conductor LIKE ? OR Tipo_Ingreso LIKE ? OR Nota LIKE ? "
            "OR CAST(Ingreso_Total AS TEXT) LIKE ?)"
        )
        comodin = f"%{texto}%"
        params_ing.extend([comodin, comodin, comodin, comodin])

        condiciones_gas.append(
            "(Categoria LIKE ? OR Comentario LIKE ? OR CAST(Monto AS TEXT) LIKE ?)"
        )
        params_gas.extend([comodin, comodin, comodin])

    sql_ing = """
        SELECT
            'Ingreso' AS Tipo,
            'ING-' || id AS ID_Visual,
            id AS ID_Numero,
            Fecha,
            Ingreso_Total AS Entrada,
            CASE
                WHEN Tipo_Ingreso = 'OTRO' THEN 'Otro ingreso'
                ELSE Conductor
            END AS Conductor,
            Parte_Triciclo AS Triciclo,
            Parte_Propietario AS Propietario,
            Parte_Chofer AS Salario,
            CASE WHEN Parte_Otro = 0 THEN NULL ELSE Parte_Otro END AS Otros,
            NULL AS Gastos,
            CASE
                WHEN Tipo_Ingreso = 'OTRO' THEN
                    CASE
                        WHEN Nota IS NULL OR Nota = '' THEN 'Otro ingreso'
                        ELSE 'Otro ingreso: ' || Nota
                    END
                WHEN Nota IS NULL OR Nota = '' THEN
                    CASE WHEN Es_Balance_Inicial = 1 THEN 'Saldo inicial' ELSE '' END
                ELSE Nota
            END AS Detalle,
            Es_Balance_Inicial AS Es_Balance
        FROM ingresos
    """
    if condiciones_ing:
        sql_ing += " WHERE " + " AND ".join(condiciones_ing)

    sql_gas = """
        SELECT
            'Gasto' AS Tipo,
            'GAS-' || id AS ID_Visual,
            id AS ID_Numero,
            Fecha,
            NULL AS Entrada,
            '' AS Conductor,
            NULL AS Triciclo,
            NULL AS Propietario,
            NULL AS Salario,
            NULL AS Otros,
            Monto AS Gastos,
            CASE
                WHEN Comentario IS NULL OR Comentario = '' THEN Categoria
                ELSE Categoria || ': ' || Comentario
            END AS Detalle,
            0 AS Es_Balance
        FROM gastos
    """
    if condiciones_gas:
        sql_gas += " WHERE " + " AND ".join(condiciones_gas)

    consultas: list[tuple[str, list[Any]]] = []
    if tipo in ("Todos", "Ingresos"):
        consultas.append((sql_ing, params_ing))
    if tipo in ("Todos", "Gastos"):
        consultas.append((sql_gas, params_gas))

    movimientos: list[dict[str, Any]] = []
    with conectar() as conn:
        for consulta, parametros in consultas:
            movimientos.extend(dict(fila) for fila in conn.execute(consulta, parametros))

    if ascendente:
        movimientos.sort(
            key=lambda x: (
                x["Fecha"],
                0 if x["Tipo"] == "Ingreso" else 1,
                x["ID_Numero"],
            )
        )
    else:
        movimientos.sort(
            key=lambda x: (x["Fecha"], x["ID_Numero"]),
            reverse=True,
        )
    if limite is not None:
        movimientos = movimientos[:limite]
    return movimientos


def obtener_registro(tipo: str, id_numero: int) -> dict[str, Any] | None:
    with conectar() as conn:
        if tipo == "Ingreso":
            fila = conn.execute(
                """
                SELECT id, Fecha, Conductor, Ingreso_Total, Nota,
                       Es_Balance_Inicial, Tipo_Ingreso
                FROM ingresos
                WHERE id = ?
                """,
                (id_numero,),
            ).fetchone()
        else:
            fila = conn.execute(
                """
                SELECT id, Fecha, Monto, Categoria, Comentario
                FROM gastos
                WHERE id = ?
                """,
                (id_numero,),
            ).fetchone()
    return dict(fila) if fila else None


def actualizar_ingreso(
    id_numero: int,
    fecha: str,
    conductor: str,
    monto: Decimal,
    nota: str,
    tipo_ingreso: str = TIPO_INGRESO_TRICICLO,
) -> None:
    actual = obtener_registro("Ingreso", id_numero)
    if not actual:
        raise ValueError("El ingreso ya no existe.")
    if actual["Es_Balance_Inicial"]:
        raise ValueError(
            "El saldo inicial se protege para evitar cambios accidentales."
        )

    if tipo_ingreso not in (TIPO_INGRESO_TRICICLO, TIPO_INGRESO_OTRO):
        raise ValueError("El tipo de ingreso no es válido.")

    monto = _validar_monto_entero(monto)
    if tipo_ingreso == TIPO_INGRESO_OTRO:
        conductor = ""
        reparto = {
            "triciclo": Decimal("0"),
            "chofer": Decimal("0"),
            "propietario": Decimal("0"),
            "otro": monto,
        }
    else:
        if not conductor.strip():
            raise ValueError("Selecciona un conductor.")
        reparto = calcular_reparto(monto)
        reparto["otro"] = Decimal("0")

    with transaccion() as conn:
        conn.execute(
            """
            UPDATE ingresos
            SET Fecha = ?, Conductor = ?, Ingreso_Total = ?,
                Parte_Triciclo = ?, Parte_Chofer = ?,
                Parte_Propietario = ?, Parte_Otro = ?, Tipo_Ingreso = ?, Nota = ?
            WHERE id = ?
            """,
            (
                fecha,
                conductor,
                decimal_a_float(monto),
                decimal_a_float(reparto["triciclo"]),
                decimal_a_float(reparto["chofer"]),
                decimal_a_float(reparto["propietario"]),
                decimal_a_float(reparto["otro"]),
                tipo_ingreso,
                nota.strip(),
                id_numero,
            ),
        )


def actualizar_gasto(
    id_numero: int,
    fecha: str,
    monto: Decimal,
    categoria: str,
    comentario: str,
) -> None:
    monto = _validar_monto_entero(monto)
    with transaccion() as conn:
        conn.execute(
            """
            UPDATE gastos
            SET Fecha = ?, Monto = ?, Categoria = ?, Comentario = ?
            WHERE id = ?
            """,
            (
                fecha,
                decimal_a_float(monto),
                categoria,
                comentario.strip(),
                id_numero,
            ),
        )


def eliminar_registro(tipo: str, id_numero: int) -> None:
    if tipo == "Ingreso":
        actual = obtener_registro(tipo, id_numero)
        if actual and actual["Es_Balance_Inicial"]:
            raise ValueError(
                "El saldo inicial no puede eliminarse desde la tabla."
            )
        tabla = "ingresos"
    else:
        tabla = "gastos"

    with transaccion() as conn:
        conn.execute(f"DELETE FROM {tabla} WHERE id = ?", (id_numero,))


def obtener_resumen_anual(anio: int) -> list[dict[str, float | int | bool]]:
    """
    Calcula el resumen de cada mes.

    Para el año en curso, los meses que todavía no han llegado se muestran
    completamente en cero. El saldo solo avanza cuando comienza realmente
    cada mes.
    """
    resumen: list[dict[str, float | int | bool]] = []
    saldo_acumulado = obtener_saldo_antes_de(f"{anio}-01-01")
    ahora = datetime.now()

    with conectar() as conn:
        for mes in range(1, 13):
            es_futuro = anio > ahora.year or (
                anio == ahora.year and mes > ahora.month
            )

            if es_futuro:
                resumen.append(
                    {
                        "mes": mes,
                        "saldo_inicial": 0.0,
                        "ingreso_bruto": 0.0,
                        "disponible": 0.0,
                        "otros": 0.0,
                        "salarios": 0.0,
                        "gastos": 0.0,
                        "saldo_final": 0.0,
                        "es_futuro": True,
                    }
                )
                continue

            patron = f"{anio}-{mes:02d}-%"
            ingreso = conn.execute(
                """
                SELECT
                    COALESCE(SUM(Ingreso_Total), 0) AS bruto,
                    COALESCE(
                        SUM(Parte_Triciclo + Parte_Propietario + Parte_Otro), 0
                    ) AS disponible,
                    COALESCE(SUM(Parte_Otro), 0) AS otros,
                    COALESCE(SUM(Parte_Chofer), 0) AS salarios
                FROM ingresos
                WHERE Fecha LIKE ?
                """,
                (patron,),
            ).fetchone()
            gasto = conn.execute(
                "SELECT COALESCE(SUM(Monto), 0) AS total FROM gastos WHERE Fecha LIKE ?",
                (patron,),
            ).fetchone()["total"]

            saldo_inicial = saldo_acumulado
            saldo_acumulado = (
                saldo_inicial
                + float(ingreso["disponible"])
                - float(gasto)
            )

            resumen.append(
                {
                    "mes": mes,
                    "saldo_inicial": saldo_inicial,
                    "ingreso_bruto": float(ingreso["bruto"]),
                    "disponible": float(ingreso["disponible"]),
                    "otros": float(ingreso["otros"]),
                    "salarios": float(ingreso["salarios"]),
                    "gastos": float(gasto),
                    "saldo_final": saldo_acumulado,
                    "es_futuro": False,
                }
            )
    return resumen

def obtener_saldo_antes_de(fecha_iso: str) -> float:
    with conectar() as conn:
        ingresos = conn.execute(
            """
            SELECT COALESCE(
                SUM(Parte_Triciclo + Parte_Propietario + Parte_Otro), 0
            ) AS total
            FROM ingresos
            WHERE Fecha < ?
            """,
            (fecha_iso,),
        ).fetchone()["total"]
        gastos = conn.execute(
            "SELECT COALESCE(SUM(Monto), 0) AS total FROM gastos WHERE Fecha < ?",
            (fecha_iso,),
        ).fetchone()["total"]
    return float(ingresos) - float(gastos)


def _validar_base_datos(ruta: Path) -> None:
    """Comprueba que el archivo sea SQLite y tenga las tablas antiguas mínimas."""
    ruta = Path(ruta)
    if not ruta.exists() or not ruta.is_file():
        raise FileNotFoundError("No se encontró el archivo de base de datos.")

    requeridas = {
        "ingresos": {
            "id", "Fecha", "Conductor", "Ingreso_Total",
            "Parte_Triciclo", "Parte_Chofer", "Parte_Propietario",
        },
        "gastos": {"id", "Fecha", "Monto", "Comentario"},
    }

    conn = sqlite3.connect(str(ruta))
    try:
        revision = conn.execute("PRAGMA quick_check").fetchone()
        if not revision or revision[0] != "ok":
            raise ValueError("El archivo SQLite está dañado o no pasó la comprobación.")

        tablas = {
            fila[0]
            for fila in conn.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table'"
            )
        }
        for tabla, columnas_necesarias in requeridas.items():
            if tabla not in tablas:
                raise ValueError(
                    f"El archivo no contiene la tabla obligatoria '{tabla}'."
                )
            columnas = {
                fila[1] for fila in conn.execute(f"PRAGMA table_info({tabla})")
            }
            faltantes = columnas_necesarias - columnas
            if faltantes:
                lista = ", ".join(sorted(faltantes))
                raise ValueError(
                    f"A la tabla '{tabla}' le faltan estas columnas: {lista}."
                )
    except sqlite3.DatabaseError as exc:
        raise ValueError("El archivo seleccionado no es una base SQLite válida.") from exc
    finally:
        conn.close()


def _copiar_sqlite(origen: Path, destino: Path) -> Path:
    """Crea una copia consistente utilizando la función backup de SQLite."""
    origen = Path(origen)
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)

    if destino.exists():
        destino.unlink()

    conn_origen = sqlite3.connect(str(origen))
    conn_destino = sqlite3.connect(str(destino))
    try:
        conn_origen.backup(conn_destino)
    finally:
        conn_destino.close()
        conn_origen.close()
    return destino


def crear_copia_seguridad() -> Path:
    if not DB_PATH.exists():
        raise FileNotFoundError("La base de datos todavía no existe.")

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    marca = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    destino = BACKUP_DIR / f"renta_triciclo_{marca}.db"
    return _copiar_sqlite(DB_PATH, destino)


def exportar_base_datos(destino: Path) -> Path:
    """Exporta una copia completa y utilizable de la base de datos actual."""
    destino = Path(destino)
    if destino.resolve() == DB_PATH.resolve():
        raise ValueError("Elige otro nombre o ubicación para la copia exportada.")
    if destino.suffix.lower() != ".db":
        destino = destino.with_suffix(".db")

    ruta = _copiar_sqlite(DB_PATH, destino)
    _validar_base_datos(ruta)
    return ruta


def importar_base_datos(origen: Path) -> Path | None:
    """
    Sustituye la base actual por otra compatible.

    Antes de hacerlo crea automáticamente una copia de seguridad de la base
    vigente. Las columnas nuevas se agregan mediante las migraciones normales.
    """
    origen = Path(origen)
    _validar_base_datos(origen)

    if origen.resolve() == DB_PATH.resolve():
        raise ValueError("La base seleccionada ya es la base que está usando el programa.")

    respaldo = crear_copia_seguridad() if DB_PATH.exists() else None
    temporal = DB_PATH.with_name(f".{DB_PATH.name}.importando")

    try:
        _copiar_sqlite(origen, temporal)
        _validar_base_datos(temporal)

        for sufijo in ("-wal", "-shm"):
            auxiliar = Path(f"{DB_PATH}{sufijo}")
            if auxiliar.exists():
                auxiliar.unlink()

        os.replace(temporal, DB_PATH)
        inicializar_base_de_datos()
    except Exception:
        if temporal.exists():
            temporal.unlink()
        raise

    return respaldo

def exportar_excel() -> Path:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    marca = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    destino = EXPORT_DIR / f"Reporte_Finanzas_Triciclo_{marca}.xlsx"

    wb = Workbook()
    ws_ing = wb.active
    ws_ing.title = "Ingresos"
    ws_gas = wb.create_sheet("Gastos")
    ws_res = wb.create_sheet("Resumen")

    encabezado_fill = PatternFill("solid", fgColor="1F4E78")
    encabezado_font = Font(color="FFFFFF", bold=True)

    def aplicar_encabezado(ws, columnas: list[str]) -> None:
        ws.append(columnas)
        for celda in ws[1]:
            celda.fill = encabezado_fill
            celda.font = encabezado_font
            celda.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    aplicar_encabezado(
        ws_ing,
        [
            "ID",
            "Fecha",
            "Conductor",
            "Tipo de ingreso",
            "Ingreso Total",
            "Parte Triciclo",
            "Parte Chofer",
            "Parte Propietario",
            "Parte Otro",
            "Balance Inicial",
            "Nota",
        ],
    )
    aplicar_encabezado(
        ws_gas,
        ["ID", "Fecha", "Monto", "Categoría", "Comentario"],
    )

    with conectar() as conn:
        for fila in conn.execute(
            """
            SELECT id, Fecha, Conductor, Tipo_Ingreso, Ingreso_Total,
                   Parte_Triciclo, Parte_Chofer, Parte_Propietario, Parte_Otro,
                   Es_Balance_Inicial, Nota
            FROM ingresos
            ORDER BY Fecha, id
            """
        ):
            ws_ing.append(list(fila))

        for fila in conn.execute(
            """
            SELECT id, Fecha, Monto, Categoria, Comentario
            FROM gastos
            ORDER BY Fecha, id
            """
        ):
            ws_gas.append(list(fila))

    # El filtro debe configurarse después de cargar los datos; de lo contrario
    # openpyxl conserva únicamente el rango del encabezado (A1:K1 / A1:E1).
    ws_ing.auto_filter.ref = ws_ing.dimensions
    ws_gas.auto_filter.ref = ws_gas.dimensions

    stats = obtener_estadisticas_globales()
    ws_res.append(["Concepto", "Monto"])
    for celda in ws_res[1]:
        celda.fill = encabezado_fill
        celda.font = encabezado_font
    ws_res.append(["Total ingresado", stats["total_ingresado"]])
    ws_res.append(["Total triciclo", stats["total_triciclo"]])
    ws_res.append(["Total propietario", stats["total_propietario"]])
    ws_res.append(["Total otros ingresos", stats["total_otros"]])
    ws_res.append(["Total choferes", stats["total_chofer"]])
    ws_res.append(["Total gastos", stats["total_gastos"]])
    ws_res.append(["Balance disponible", stats["balance"]])

    for ws in (ws_ing, ws_gas, ws_res):
        for columna in ws.columns:
            ancho = max(len(str(celda.value or "")) for celda in columna) + 2
            ws.column_dimensions[get_column_letter(columna[0].column)].width = min(
                max(ancho, 12), 35
            )

    for ws in (ws_ing, ws_gas):
        for fila in ws.iter_rows(min_row=2):
            for celda in fila:
                if isinstance(celda.value, (int, float)) and celda.column > 2:
                    celda.number_format = '$#,##0.00'

    for celda in ws_res["B"][1:]:
        celda.number_format = '$#,##0.00'

    wb.save(destino)
    return destino
